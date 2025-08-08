import openpyxl
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

wb = openpyxl.load_workbook('students_table.xlsx')
wb_sheet = wb.worksheets[0]
values = Reference(wb_sheet, min_col=2, min_row=5,
                   max_col=6, max_row=5)

titles = Reference(wb_sheet, min_col=2, min_row=1,
                   max_row=1, max_col=6)
chart = BarChart()
chart.add_data(values, titles)
chart.set_categories(titles)
wb_sheet.add_chart(chart,"H7")
chart.style = 8 # add bar color
chart.x_axis.title = 'Students' # add label in the x axes title
chart.y_axis.title = 'Scores'
chart.legend.position = 'tr'
chart.title = "Students's scores average"
wb.save("students_table.xlsx")
